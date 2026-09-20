"""MDF to Excel Converter.

Reads a Microsoft SQL Server .mdf data file directly (no SQL Server needed)
and exports every user table into a single Excel (.xlsx) workbook, with one
sheet per table plus a "_Summary" sheet.

Supported SQL Server file versions: 2000 (539), 2005 (611/612), 2008 (655),
2008 R2 (661).

Usage (command line):
    python mdf2excel.py "database.mdf"
    python mdf2excel.py "database.mdf" -o "output.xlsx"
    python mdf2excel.py "database.mdf" --tables
    python mdf2excel.py "database.mdf" --info

With no arguments the tool opens a small graphical window.
"""

import os
import re
import sys
import queue
import threading
import traceback
from datetime import datetime, date

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font
except ImportError:  # pragma: no cover
    Workbook = None
    Font = None

import mdf_parser


INVALID_SHEET_CHARS = re.compile(r'[\[\]\*/\\?:]')
ILLEGAL_XML_CHARS = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f]')
CELL_LIMIT = 32767


def clean_string(s):
    return ILLEGAL_XML_CHARS.sub('', s)


def make_sheet_name(name, used):
    n = INVALID_SHEET_CHARS.sub('_', name)
    n = n.replace("'", '_').replace('"', '_').strip()
    if not n:
        n = 'Sheet'
    if len(n) > 31:
        n = n[:31].rstrip()
    base = n
    idx = 2
    while n.lower() in used:
        suffix = f'_{idx}'
        n = base[:31 - len(suffix)] + suffix
        idx += 1
    used.add(n.lower())
    return n


def to_cell(v):
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.strftime('%Y-%m-%d %H:%M:%S')
    if isinstance(v, date):
        return v.strftime('%Y-%m-%d')
    if isinstance(v, bool):
        return int(v)
    if isinstance(v, int):
        return v
    if isinstance(v, float):
        if v != v or v in (float('inf'), float('-inf')):
            return str(v)
        return v
    if isinstance(v, bytes):
        return v.hex()
    s = clean_string(str(v))
    return s if len(s) <= CELL_LIMIT else s[:CELL_LIMIT]


def apply_decimal_scale(row, schema):
    for col in schema:
        if col.xtype in (mdf_parser.XTYPE_DECIMAL, mdf_parser.XTYPE_NUMERIC) and col.scale > 0:
            v = row.get(col.name)
            if isinstance(v, int) and not isinstance(v, bool):
                row[col.name] = v / (10 ** col.scale)


def convert_mdf(mdf_path, out_path, progress=None):
    """Convert an MDF file to a single Excel workbook.

    progress(percent:int, message:str) is called periodically (optional).
    Returns (out_path, summary) where summary is a list of
    (table_name, row_count, column_count, status) tuples.
    """
    if Workbook is None:
        raise RuntimeError('openpyxl is not installed. Install it with: pip install openpyxl')

    summary = []
    used_names = set()
    wb = Workbook()

    with mdf_parser.MdfFile(mdf_path) as mdf:
        db_name = mdf.get_db_name()
        version = mdf.get_version()
        catalog = mdf_parser.SystemCatalog(mdf)
        reader = mdf_parser.TableReader(mdf, catalog)
        tables = catalog.get_table_list()
        tables.sort(key=lambda t: t[1].lower())
        total = len(tables)

        first = True
        for i, (tid, tname) in enumerate(tables):
            schema = catalog.get_table_schema(tid)
            pct = int((i / total) * 100) if total else 0
            if progress:
                progress(pct, f'Processing {tname} ({i + 1}/{total})')

            if not schema:
                summary.append((tname, 0, 0, 'no schema'))
                continue

            sheet_name = make_sheet_name(tname, used_names)
            if first:
                ws = wb.active
                ws.title = sheet_name
                first = False
            else:
                ws = wb.create_sheet(sheet_name)

            ws.append([col.name for col in schema])
            for c in range(1, len(schema) + 1):
                ws.cell(row=1, column=c).font = Font(bold=True)

            count = 0
            skipped = 0
            status = 'ok'
            try:
                for row in reader.iter_table(tname):
                    apply_decimal_scale(row, schema)
                    values = [to_cell(row.get(col.name)) for col in schema]
                    try:
                        ws.append(values)
                    except Exception:
                        values = [clean_string(str(v))[:CELL_LIMIT] if v is not None else None for v in values]
                        ws.append(values)
                        skipped += 1
                    count += 1
            except Exception as e:
                status = f'error: {e}'
            if skipped:
                status = f'{status} ({skipped} rows sanitized)'

            summary.append((tname, count, len(schema), status))

    summary_sheet = wb.create_sheet('_Summary', 0)
    summary_sheet.append(['Table', 'Rows', 'Columns', 'Status'])
    for c in range(1, 5):
        summary_sheet.cell(row=1, column=c).font = Font(bold=True)
    for tname, nrows, ncols, status in summary:
        summary_sheet.append([tname, nrows, ncols, clean_string(status)])

    for ws in wb.worksheets:
        if ws.title != '_Summary':
            ws.freeze_panes = 'A2'

    wb.save(out_path)
    if progress:
        progress(100, 'Done')

    return out_path, summary


def _show_db_info(mdf_path):
    with mdf_parser.MdfFile(mdf_path) as mdf:
        catalog = mdf_parser.SystemCatalog(mdf)
        print(f'DB Name: {mdf.get_db_name()}')
        print(f'Version: {mdf.get_version()}')
        print(f'Size: {mdf.size:,} bytes')
        print(f'Pages: {mdf.page_count:,}')


def _show_tables(mdf_path):
    with mdf_parser.MdfFile(mdf_path) as mdf:
        catalog = mdf_parser.SystemCatalog(mdf)
        reader = mdf_parser.TableReader(mdf, catalog)
        tables = catalog.get_table_list()
        tables.sort(key=lambda t: t[1].lower())
        print(f'{len(tables)} tables:')
        for tid, tname in tables:
            n = reader.count_rows(tname)
            print(f'  {tname}: {n} rows')


def run_cli(argv):
    import argparse
    parser = argparse.ArgumentParser(
        prog='mdf2excel',
        description='Convert a SQL Server .mdf file to a single Excel (.xlsx) workbook.')
    parser.add_argument('mdf', nargs='?', help='Path to the .mdf file')
    parser.add_argument('-o', '--output', help='Output .xlsx path')
    parser.add_argument('--info', action='store_true', help='Only show database info')
    parser.add_argument('--tables', action='store_true', help='Only list tables')
    args = parser.parse_args(argv)

    if not args.mdf:
        parser.print_help()
        return 1

    mdf_path = args.mdf
    if not os.path.exists(mdf_path):
        print(f'File not found: {mdf_path}', file=sys.stderr)
        return 1

    if args.info:
        _show_db_info(mdf_path)
        return 0
    if args.tables:
        _show_tables(mdf_path)
        return 0

    if args.output:
        out_path = args.output
    else:
        out_path = os.path.splitext(mdf_path)[0] + '_AllTables.xlsx'

    def progress(pct, msg):
        print(f'[{pct:3d}%] {msg}')

    out_path, summary = convert_mdf(mdf_path, out_path, progress)
    print(f'Saved: {out_path}')
    print(f'Tables: {len(summary)}')
    return 0


def run_gui(initial_file=None):
    try:
        import tkinter as tk
        from tkinter import filedialog, messagebox, scrolledtext, ttk
    except ImportError as e:
        print('Graphical window is not available:', e)
        print('Linux: install Tk (e.g. "sudo apt install python3-tk").')
        print('macOS: use a Python build that includes Tkinter.')
        print('You can still use the command line: python mdf2excel.py file.mdf')
        return 1

    root = tk.Tk()
    root.title('MDF to Excel Converter')
    root.geometry('760x560')
    root.minsize(620, 460)

    mdf_var = tk.StringVar(value=initial_file or '')
    out_var = tk.StringVar()

    def pick_mdf():
        f = filedialog.askopenfilename(
            title='Select MDF file',
            filetypes=[('SQL Server data files', '*.mdf'), ('All files', '*.*')])
        if f:
            mdf_var.set(f)
            if not out_var.get().strip():
                out_var.set(os.path.splitext(f)[0] + '_AllTables.xlsx')

    def pick_out():
        initial = os.path.basename(out_var.get().strip()) or 'output.xlsx'
        f = filedialog.asksaveasfilename(
            title='Save Excel as',
            defaultextension='.xlsx',
            filetypes=[('Excel workbook', '*.xlsx')],
            initialfile=initial)
        if f:
            out_var.set(f)

    frame = ttk.Frame(root, padding=10)
    frame.pack(fill='both', expand=True)

    ttk.Label(frame, text='MDF file:').grid(row=0, column=0, sticky='w')
    ttk.Entry(frame, textvariable=mdf_var).grid(row=0, column=1, sticky='ew', padx=5)
    ttk.Button(frame, text='Browse...', command=pick_mdf).grid(row=0, column=2)

    ttk.Label(frame, text='Output Excel:').grid(row=1, column=0, sticky='w', pady=(6, 0))
    ttk.Entry(frame, textvariable=out_var).grid(row=1, column=1, sticky='ew', padx=5, pady=(6, 0))
    ttk.Button(frame, text='Browse...', command=pick_out).grid(row=1, column=2, pady=(6, 0))

    frame.columnconfigure(1, weight=1)

    btn_row = ttk.Frame(frame)
    btn_row.grid(row=2, column=0, columnspan=3, sticky='ew', pady=10)
    convert_btn = ttk.Button(btn_row, text='Convert to Excel')
    convert_btn.pack(side='left')
    status_var = tk.StringVar(value='Ready')
    ttk.Label(btn_row, textvariable=status_var).pack(side='left', padx=10)

    progress_bar = ttk.Progressbar(frame, orient='horizontal', mode='determinate', maximum=100)
    progress_bar.grid(row=3, column=0, columnspan=3, sticky='ew')

    log_text = scrolledtext.ScrolledText(frame, height=20, state='disabled', wrap='word')
    log_text.grid(row=4, column=0, columnspan=3, sticky='nsew', pady=(10, 0))
    frame.rowconfigure(4, weight=1)

    q = queue.Queue()

    def append_log(msg):
        log_text.configure(state='normal')
        log_text.insert('end', msg)
        log_text.see('end')
        log_text.configure(state='disabled')

    def worker(mdf_path, out_path):
        try:
            def progress(pct, msg):
                q.put(('progress', pct, msg))
            out, summary = convert_mdf(mdf_path, out_path, progress)
            q.put(('done', out, summary))
        except Exception:
            q.put(('error', traceback.format_exc()))

    def poll():
        try:
            while True:
                evt = q.get_nowait()
                kind = evt[0]
                if kind == 'progress':
                    _, pct, msg = evt
                    progress_bar['value'] = pct
                    status_var.set(msg)
                    append_log(msg + '\n')
                elif kind == 'done':
                    _, out, summary = evt
                    progress_bar['value'] = 100
                    status_var.set('Done')
                    append_log('\nSaved: ' + out + '\n')
                    for tname, nrows, ncols, st in summary:
                        append_log(f'{tname}: {nrows} rows, {ncols} cols, {st}\n')
                    messagebox.showinfo('Done', f'Converted {len(summary)} tables.\n\n{out}')
                    convert_btn.config(state='normal')
                    return
                elif kind == 'error':
                    _, tb = evt
                    append_log('ERROR\n' + tb + '\n')
                    messagebox.showerror('Error', tb)
                    convert_btn.config(state='normal')
                    return
        except queue.Empty:
            pass
        root.after(100, poll)

    def start():
        mdf_path = mdf_var.get().strip()
        out_path = out_var.get().strip()
        if not mdf_path:
            messagebox.showerror('Error', 'Please choose a .mdf file.')
            return
        if not os.path.exists(mdf_path):
            messagebox.showerror('Error', 'MDF file not found:\n' + mdf_path)
            return
        if not out_path:
            out_path = os.path.splitext(mdf_path)[0] + '_AllTables.xlsx'
            out_var.set(out_path)

        convert_btn.config(state='disabled')
        progress_bar['value'] = 0
        status_var.set('Starting...')
        log_text.configure(state='normal')
        log_text.delete('1.0', 'end')
        log_text.configure(state='disabled')

        t = threading.Thread(target=worker, args=(mdf_path, out_path), daemon=True)
        t.start()
        root.after(100, poll)

    convert_btn.config(command=start)
    root.mainloop()
    return 0


def main(argv=None):
    if sys.stdout is None:
        sys.stdout = open(os.devnull, 'w', encoding='utf-8')
    if sys.stderr is None:
        sys.stderr = open(os.devnull, 'w', encoding='utf-8')
    if hasattr(sys.stdout, 'reconfigure'):
        sys.stdout.reconfigure(encoding='utf-8', errors='replace')
    if hasattr(sys.stderr, 'reconfigure'):
        sys.stderr.reconfigure(encoding='utf-8', errors='replace')

    if argv is None:
        argv = sys.argv[1:]

    cli_flags = {'--cli', '--info', '--tables', '--help', '-h', '-o', '--output'}

    # Force CLI mode when an MDF path is combined with explicit flags.
    if any(a in cli_flags for a in argv):
        return run_cli([a for a in argv if a != '--cli'])

    if argv and argv[0].lower().endswith('.mdf') and os.path.exists(argv[0]):
        return run_cli(argv)

    # GUI mode; prefill the first .mdf path passed via drag-and-drop, if any.
    initial = None
    for a in argv:
        if a.lower().endswith('.mdf') and os.path.exists(a):
            initial = a
            break
    return run_gui(initial)


if __name__ == '__main__':
    sys.exit(main())
